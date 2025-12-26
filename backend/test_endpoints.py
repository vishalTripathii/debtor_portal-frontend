import requests
import json

# Test Bulk Delete Endpoint
print("=" * 80)
print("TESTING BULK DELETE ENDPOINT")
print("=" * 80)

url = "https://jxnu8wrip2.execute-api.ap-southeast-1.amazonaws.com/dev/bulk/delete-excel"

# Create test Excel file with account numbers
import openpyxl
wb = openpyxl.Workbook()
ws = wb.active
ws['A1'] = 'Account Number'
test_accounts = [
    '202203000242631',
    '201912000003482',
    '202003000018467'
]
for idx, acc in enumerate(test_accounts, start=2):
    ws[f'A{idx}'] = acc

wb.save('test_delete_temp.xlsx')

# Test 1: Send POST without file (should fail)
print("\n[TEST 1] POST without file (expected to fail):")
print("-" * 80)
response = requests.post(url)
print(f"Status Code: {response.status_code}")
print(f"Response: {response.text[:500]}")

# Test 2: Send POST with file but no auth token (should fail or succeed based on auth)
print("\n[TEST 2] POST with Excel file:")
print("-" * 80)
files = {'file': open('test_delete_temp.xlsx', 'rb')}
response = requests.post(url, files=files)
print(f"Status Code: {response.status_code}")
print(f"Response: {response.text[:500]}")

# Test 3: Check if API Gateway is reachable
print("\n[TEST 3] OPTIONS request (CORS preflight):")
print("-" * 80)
response = requests.options(url)
print(f"Status Code: {response.status_code}")
print(f"Headers: {dict(response.headers)}")

print("\n" + "=" * 80)
print("TEST COMPLETE")
print("=" * 80)
