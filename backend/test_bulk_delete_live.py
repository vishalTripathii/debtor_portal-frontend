import requests
import os
import openpyxl
from datetime import datetime

print("="*80)
print("BULK DELETE - DEPLOYMENT TEST")
print("="*80)
print(f"Test Time: {datetime.now()}")
print("="*80)

# Configuration
EXCEL_FILE = r"C:\Users\tripa\Downloads\debtor_upload_template (3).xlsx"
API_ENDPOINT = "https://jxnu8wrip2.execute-api.ap-southeast-1.amazonaws.com/dev/bulk/delete-excel"

# Step 1: Check if file exists
print("\n[STEP 1] Checking Excel file...")
if os.path.exists(EXCEL_FILE):
    print(f"✓ File found: {EXCEL_FILE}")
    file_size = os.path.getsize(EXCEL_FILE) / 1024  # KB
    print(f"  File size: {file_size:.2f} KB")
else:
    print(f"✗ File NOT found: {EXCEL_FILE}")
    exit(1)

# Step 2: Read and analyze Excel content
print("\n[STEP 2] Analyzing Excel content...")
try:
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb.active
    
    # Find Account Number column
    headers = []
    for cell in ws[1]:
        headers.append(cell.value)
    
    print(f"  Columns found: {headers}")
    
    if 'Account Number' not in headers:
        print("✗ 'Account Number' column not found!")
        exit(1)
    
    # Count account numbers
    account_col_idx = headers.index('Account Number') + 1
    account_numbers = []
    
    for row in range(2, ws.max_row + 1):
        account = ws.cell(row=row, column=account_col_idx).value
        if account:
            account_numbers.append(str(account).strip())
    
    print(f"✓ Found {len(account_numbers)} account numbers")
    print(f"  First 5: {account_numbers[:5]}")
    print(f"  Last 5: {account_numbers[-5:]}")
    
except Exception as e:
    print(f"✗ Error reading Excel: {e}")
    exit(1)

# Step 3: Test API endpoint connectivity
print("\n[STEP 3] Testing API endpoint connectivity...")
try:
    response = requests.get(API_ENDPOINT.replace('/bulk/delete-excel', '/'))
    print(f"  API Gateway is reachable")
except Exception as e:
    print(f"✗ Cannot reach API Gateway: {e}")

# Step 4: Send bulk delete request
print("\n[STEP 4] Sending bulk delete request...")
print(f"  Endpoint: {API_ENDPOINT}")
print(f"  File: {os.path.basename(EXCEL_FILE)}")
print(f"  Accounts to delete: {len(account_numbers)}")
print("\n  Sending request... (this may take a while)")

try:
    with open(EXCEL_FILE, 'rb') as f:
        files = {'file': (os.path.basename(EXCEL_FILE), f, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
        response = requests.post(API_ENDPOINT, files=files, timeout=300)
    
    print(f"\n✓ Response received!")
    print(f"  Status Code: {response.status_code}")
    print(f"  Response Time: {response.elapsed.total_seconds():.2f} seconds")
    
    # Step 5: Analyze response
    print("\n[STEP 5] Analyzing response...")
    print("-"*80)
    
    if response.status_code == 200:
        try:
            result = response.json()
            print("✓ SUCCESS - Bulk delete completed!")
            print(f"\n  Results:")
            print(f"    - Total in Excel: {result.get('total_in_excel', 'N/A')}")
            print(f"    - Successfully deleted: {result.get('deleted_count', 'N/A')}")
            print(f"    - Not found (already deleted): {result.get('not_found_count', 'N/A')}")
            print(f"    - Message: {result.get('message', 'N/A')}")
            
            if result.get('success'):
                deletion_rate = (result.get('deleted_count', 0) / len(account_numbers)) * 100
                print(f"\n  Deletion Rate: {deletion_rate:.1f}%")
        except Exception as e:
            print(f"  Response: {response.text[:500]}")
    else:
        print(f"✗ FAILED - Status {response.status_code}")
        print(f"  Response: {response.text[:500]}")
        
        if response.status_code == 401:
            print("\n  Note: 401 means authentication is required")
            print("  This is expected if testing directly without token")
        elif response.status_code == 500:
            print("\n  Note: 500 means server error")
            print("  Check CloudWatch logs for details")

except requests.exceptions.Timeout:
    print("✗ Request timed out (>300 seconds)")
except Exception as e:
    print(f"✗ Error: {e}")

# Step 6: Summary
print("\n"+"="*80)
print("TEST SUMMARY")
print("="*80)
print(f"File: {os.path.basename(EXCEL_FILE)}")
print(f"Total Account Numbers: {len(account_numbers)}")
print(f"Endpoint: {API_ENDPOINT}")
print(f"Status: Check response above")
print("="*80)

# Step 7: Next steps
print("\n[NEXT STEPS]")
print("1. If 401 error: Need to add authentication token")
print("2. If 200 success: Check MongoDB to verify deletions")
print("3. View logs: aws logs tail /aws/lambda/debtor-portal-bulk-operations-dev --follow")
print("4. Test from frontend: https://d1hmzuewg6k3ss.cloudfront.net")
